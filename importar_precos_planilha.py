# -*- coding: utf-8 -*-
"""
Leva os precos da planilha "Cadastro de Precos - <Mes> <Ano>.xlsx" para dentro
do sistema, de uma vez.

Existe para a virada: o preco passou a ser cadastrado na tela de Movimentacoes,
e a planilha vai ser descontinuada. Este script e a ponte — traz o que ja foi
digitado nela para nao comecar do zero na tela.

Onde o preco vai morar: dados/lancamentos.json, em
    [data][casa]["insumos"][chave]["precos"][tipo]
com tipo = "entrada" (aba 2 da planilha) ou "transferencia" (aba 3). E a mesma
granularidade da planilha — um valor por data, casa, item e tipo de movimento —,
entao a importacao e linha a linha, sem media nem rateio.

O que NAO entra, e por que:

  * Itens que o CDE nao acompanha: raspa, apara e pele de salmao e camarao P
    c/rabo. Eles existem na planilha da casa e na de precos, mas nao sao insumo
    do sistema (veja a lista IGNORAR do importar_diario.py). O preco deles fica
    sem destino e e listado no relatorio.
  * Linhas cujo movimento nao existe no sistema. O preco se prende a um
    movimento; sem a chegada lancada, nao ha onde pendurar. Acontece com os
    itens IN NATURA se a sincronizacao ainda nao rodou depois da mudanca que
    passou a importa-los — nesse caso rode primeiro:
        python importar_diario.py --gravar

Uso:
    python importar_precos_planilha.py                 mostra o que faria
    python importar_precos_planilha.py --gravar        grava
    python importar_precos_planilha.py --mes 08        outro mes
"""
import json
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

RAIZ = Path(__file__).parent
DESTINO = RAIZ / "dados" / "lancamentos.json"
BACKUPS = RAIZ / "backups"

MESES = ["", "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
         "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]

# O nome das abas carrega o mes; ficava fixo em 'agosto'.
def abas_do_mes(mes):
    m = MESES[mes].lower()
    return [("2. Entradas de %s" % m, "entrada"),
            ("3. Transferências de %s" % m, "transferencia")]

# Nome do item na planilha -> chave do insumo no CDE.
# None = a planilha tem preco para isso, mas o sistema nao acompanha o item.
DE_PARA = {
    "Salmão — libra 08/10": "salmao_0810",
    "Salmão — libra 10/12": "salmao_1012",
    "Salmão — libra 12/14": "salmao_1214",
    "Salmão — libra 14/16": "salmao_1416",
    "Salmão — raspa": None,
    "Salmão — apara": None,
    "Salmão — pele": None,
    "Camarão G s/rabo": "camarao_g",
    "Camarão M s/rabo": "camarao_m",
    "Camarão P s/rabo": "camarao_p",
    "Camarão P c/rabo": None,
    "Kani": "kani",
    "Anchova": "anchova",
    "Nori": "nori",
    "Cream cheese": "cream_cheese",
    "Arroz de sushi": "arroz",
    "Filé de tilápia": "file_tilapia",
    "Anel de lula — in natura": "lula_in_natura",
    "Anel de lula — processado": "lula",
    "Tentáculo de lula — in natura": "polvo_in_natura",
    "Tentáculo de lula — processado": "polvo",
    "Lombo de atum — in natura": "atum_in_natura",
    "Lombo de atum — processado": "atum",
    "Patinho": "patinho",
}


def ler_compartilhado(arq):
    """Bytes do arquivo mesmo com o Excel ou o OneDrive segurando ele."""
    ps = (
        "$fs=New-Object System.IO.FileStream("
        "'%s',[System.IO.FileMode]::Open,[System.IO.FileAccess]::Read,"
        "([System.IO.FileShare]::ReadWrite -bor [System.IO.FileShare]::Delete));"
        "$ms=New-Object System.IO.MemoryStream;$fs.CopyTo($ms);$fs.Close();"
        "[Console]::OpenStandardOutput().Write($ms.ToArray(),0,$ms.Length)"
    ) % str(arq).replace("'", "''")
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                           capture_output=True, timeout=180)
    except (OSError, subprocess.TimeoutExpired):
        return None
    return r.stdout if r.returncode == 0 and r.stdout else None


def abrir_dois(arq):
    """A mesma planilha lida de dois jeitos: por valor e por formula.

    Precisa dos dois porque o openpyxl nao calcula formula. Quando o
    gerar_cadastro_precos reescreve a planilha para incluir os dias novos, o
    valor que o Excel tinha guardado em cache se perde, e a leitura por valor
    devolve None na celula onde o financeiro digitou uma conta (=14932/48).
    Sem comparar com a leitura por formula, essas linhas eram puladas caladas,
    como se estivessem em branco.
    """
    import io as _io
    import openpyxl
    try:
        with open(arq, "rb") as f:
            dados = f.read()
    except PermissionError:
        dados = ler_compartilhado(arq)
    if not dados:
        return None, None
    return (openpyxl.load_workbook(_io.BytesIO(dados), data_only=True),
            openpyxl.load_workbook(_io.BytesIO(dados), data_only=False))


def ler_planilha(arq, mes):
    """([(tipo, casa, data, item, chave, qtd, preco)], formulas sem valor)."""
    valores, formulas = abrir_dois(arq)
    if valores is None:
        raise SystemExit("Nao consegui abrir a planilha (travada?):\n  %s" % arq)
    linhas, sem_calculo = [], []
    for nome_aba, tipo in abas_do_mes(mes):
        if nome_aba not in valores.sheetnames:
            print("  [aviso] aba '%s' nao existe na planilha" % nome_aba)
            continue
        ws = valores[nome_aba]
        wf = formulas[nome_aba] if nome_aba in formulas.sheetnames else None
        for r in range(5, ws.max_row + 1):
            data, casa = ws.cell(r, 1).value, ws.cell(r, 2).value
            item, qtd, preco = (ws.cell(r, 4).value, ws.cell(r, 6).value,
                                ws.cell(r, 7).value)
            if data is None or not hasattr(data, "strftime") or not item:
                continue
            if preco in (None, ""):
                formula = wf.cell(r, 7).value if wf is not None else None
                if isinstance(formula, str) and formula.startswith("="):
                    sem_calculo.append((data.strftime("%Y-%m-%d"), casa, item,
                                        formula))
                continue                      # nao preenchido: nada a importar
            linhas.append((tipo, casa, data.strftime("%Y-%m-%d"), item,
                           DE_PARA.get(item, "?"), qtd, float(preco)))
    return linhas, sem_calculo


def main():
    args = sys.argv[1:]
    gravar = "--gravar" in args
    hoje = datetime.now()
    mes, ano = hoje.month, hoje.year
    if "--mes" in args:
        mes = int(args[args.index("--mes") + 1])

    arq = (RAIZ / ".." / ("Cadastro de Precos - %s %d.xlsx" % (MESES[mes], ano))).resolve()
    print("planilha: %s" % arq)
    if not arq.exists():
        raise SystemExit("planilha nao encontrada")
    if not DESTINO.exists():
        raise SystemExit("dados/lancamentos.json nao encontrado")

    linhas, sem_calculo = ler_planilha(arq, mes)
    with open(DESTINO, encoding="utf-8") as f:
        dados = json.load(f)

    aplicados, sem_item, sem_movimento, mudou = [], [], [], []
    for tipo, casa, iso, item, chave, qtd, preco in linhas:
        if chave is None or chave == "?":
            sem_item.append((iso, casa, item, preco, chave == "?"))
            continue
        campos = (((dados.get(iso) or {}).get(casa) or {}).get("insumos") or {}).get(chave)
        if not isinstance(campos, dict):
            sem_movimento.append((iso, casa, item, chave, tipo, "sem lançamento"))
            continue
        try:
            movimento = float(campos.get(tipo) or 0)
        except (TypeError, ValueError):
            movimento = 0.0
        if abs(movimento) <= 0.0005:
            sem_movimento.append((iso, casa, item, chave, tipo, "sem %s no dia" % tipo))
            continue
        antigo = (campos.get("precos") or {}).get(tipo)
        # Comparar com o MESMO arredondamento que se grava. Contra o float cru
        # da planilha, os 6 decimais guardados sempre "diferiam" na setima casa
        # e toda reimportacao anunciava 39 precos mudados, todos iguais.
        novo = round(preco, 6)
        if antigo is not None and abs(float(antigo) - novo) > 1e-9:
            mudou.append((iso, casa, item, tipo, float(antigo), novo))
        if gravar:
            campos.setdefault("precos", {})[tipo] = novo
        aplicados.append((iso, casa, item, tipo, preco, movimento))

    print()
    print("linhas com preco na planilha : %d" % len(linhas))
    if sem_calculo:
        print("formula SEM valor calculado  : %d (nao entram — veja abaixo)"
              % len(sem_calculo))
    print("precos que entram no sistema : %d" % len(aplicados))
    print("itens que o CDE nao acompanha: %d" % len(sem_item))
    print("sem movimento correspondente : %d" % len(sem_movimento))
    if mudou:
        print("ja tinham preco DIFERENTE    : %d (a planilha prevalece)" % len(mudou))

    if sem_item:
        print("\n--- ITENS SEM DESTINO NO CDE (preco fica de fora) ---")
        vistos = {}
        for iso, casa, item, preco, desconhecido in sem_item:
            vistos.setdefault(item, [0, desconhecido])
            vistos[item][0] += 1
        for item, (n, desconhecido) in sorted(vistos.items()):
            print("   %-34s %3d linha(s)%s" % (item, n,
                  "   <== NOME NAO MAPEADO" if desconhecido else ""))

    if sem_movimento:
        print("\n--- SEM MOVIMENTO NO SISTEMA (primeiras 15) ---")
        for iso, casa, item, chave, tipo, motivo in sem_movimento[:15]:
            print("   %s %-5s %-34s %-14s %s" % (iso, casa, item, tipo, motivo))
        if len(sem_movimento) > 15:
            print("   ... e mais %d" % (len(sem_movimento) - 15))
        proc = [s for s in sem_movimento
                if s[3] in ("lula", "polvo", "atum") or s[3].endswith("_in_natura")]
        if proc:
            print("\n   %d delas sao itens de PROCESSAMENTO (in natura / processado)."
                  "\n   O importador das planilhas mudou: antes ele descartava o lado in"
                  "\n   natura e lia o processado com as colunas deslocadas. Rode uma"
                  "\n   sincronizacao nova ANTES desta importacao e esse numero deve cair:"
                  "\n       python importar_diario.py --gravar" % len(proc))

    if mudou:
        print("\n--- PRECO QUE MUDOU (primeiros 15) ---")
        for iso, casa, item, tipo, antes, agora in mudou[:15]:
            print("   %s %-5s %-30s %-14s %10.4f -> %10.4f"
                  % (iso, casa, item, tipo, antes, agora))

    if sem_calculo:
        print("\n--- FORMULA SEM VALOR CALCULADO ---")
        print("   O openpyxl nao calcula formula, e o valor que o Excel tinha em"
              "\n   cache se perdeu quando a planilha foi reescrita. A formula esta"
              "\n   la; falta o resultado. Abra a planilha no Excel, salve (Ctrl+S)"
              "\n   e rode esta importacao de novo — as %d entram." % len(sem_calculo))
        for iso, casa, item, formula in sem_calculo:
            print("   %s %-5s %-32s %s" % (iso, casa, item, formula))

    if not gravar:
        print("\n(nada foi gravado — rode com --gravar)")
        return 0

    BACKUPS.mkdir(exist_ok=True)
    carimbo = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    shutil.copy(DESTINO, BACKUPS / ("lancamentos_antes-dos-precos_%s.json" % carimbo))
    tmp = DESTINO.with_suffix(".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=1)
    tmp.replace(DESTINO)
    print("\ngravado. backup em backups/lancamentos_antes-dos-precos_%s.json" % carimbo)
    return 0


if __name__ == "__main__":
    sys.exit(main())
