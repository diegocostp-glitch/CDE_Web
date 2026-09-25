# -*- coding: utf-8 -*-
"""
Gera/atualiza a planilha "Cadastro de Precos - <Mes> <Ano>.xlsx".

Por que existe: o painel precisa do custo unitario de cada insumo (veja o
_nota em dados/custos.json), e esse numero nao esta em lugar nenhum do
sistema — quem sabe o preco pago e o financeiro. Entao o CDE monta a lista de
TUDO que chegou e de tudo que foi transferido no mes, deixa a coluna de preco
em branco (laranja) e devolve a planilha para ser preenchida.

De onde vem cada aba:

  1. Saldo inicial   — o fechamento do mes anterior. Montada uma vez, pelo
                       --criar, e nunca mais mexida: e ali que o financeiro
                       digita o valor unitario do estoque que virou o mes.
  2. Entradas        — toda "Chegada" das abas diarias das planilhas-fonte.
  3. Transferencias  — toda "Transferencia" das mesmas abas.

IMPORTANTE — por que le a planilha-fonte e nao dados/lancamentos.json:
o lancamentos.json e um recorte. O importar_diario.py joga fora, de proposito,
o que o painel nao acompanha (veja a lista IGNORAR de la): apara, pele e raspa
de salmao, camarao P c/rabo, e o lado IN NATURA dos itens que a casa processa
(anel de lula, tentaculo de lula, lombo de atum) — deste ultimo ele guarda so o
PROCESSADO. Mas o financeiro paga pelo IN NATURA e pelos subprodutos tambem,
entao o Cadastro de Precos precisa das duas coisas. Gerar esta planilha a
partir do lancamentos.json produzia 73 linhas onde o correto sao 151, e ainda
etiquetava o processado como se fosse in natura.

O preco JA DIGITADO nunca e perdido: antes de reescrever, o script guarda o que
esta preenchido pela chave (casa, data, item) e devolve linha por linha. Assim
pode rodar quantas vezes quiser no meio do mes, para pegar os dias novos.

Uso:
    python gerar_cadastro_precos.py                 mostra o que falta, nao grava
    python gerar_cadastro_precos.py --gravar        atualiza a planilha
    python gerar_cadastro_precos.py --mes 08        outro mes (padrao: o atual)
    python gerar_cadastro_precos.py --criar         previa da planilha de um mes
                                                    que ainda nao existe
    python gerar_cadastro_precos.py --criar --gravar  monta ela
"""
import re
import shutil
import sys
import time
import unicodedata
from datetime import datetime, timedelta
from pathlib import Path

RAIZ = Path(__file__).parent
ORIGEM = (RAIZ / ".." / ".." / ".." / "Documentos 2026" /
          "CDE - Controle Diário de Estoque").resolve()
BACKUPS = RAIZ / "backups"

MESES = ["", "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
         "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]

CASAS = [("SL", "Senador Lemos"), ("DC", "Duque de Caxias"),
         ("DLU", "Delivery Umarizal"), ("DLCN", "Delivery Cidade Nova")]

COLUNAS = [2, 5, 8, 11]                  # B, E, H, K — os quatro blocos por linha
CAMPOS = [("INICIO", "inicial"), ("FINAL", "final"), ("CHEGADA", "entrada"),
          ("TRANSFERENCIA", "transferencia"), ("PERDA", "desperdicio")]

# Titulo do bloco -> nome do item na planilha de precos.
SIMPLES = [
    ("LIBRA (08/10", "Salmão — libra 08/10"),
    ("LIBRA (10/12", "Salmão — libra 10/12"),
    ("LIBRA (12/14", "Salmão — libra 12/14"),
    ("LIBRA (14/16", "Salmão — libra 14/16"),
    ("LIBRA (00/00", "Salmão — libra 14/16"),   # faixa ainda nao nomeada na fonte
    ("RASPA DE SALMAO", "Salmão — raspa"),
    ("APARA DE SALMAO", "Salmão — apara"),
    ("PELE DE SALMAO", "Salmão — pele"),
    ("CAMARAO G - S/RABO", "Camarão G s/rabo"),
    ("CAMARAO M - S/RABO", "Camarão M s/rabo"),
    ("CAMARAO P - S/RABO", "Camarão P s/rabo"),
    ("CAMARAO P - C/RABO", "Camarão P c/rabo"),
    ("KANI", "Kani"),
    ("ANCHOVA", "Anchova"),
    ("NORI", "Nori"),
    ("CREAM CHEESE", "Cream cheese"),
    ("ARROZ DE SUSHI", "Arroz de sushi"),
    ("FILE DE TILAPIA", "Filé de tilápia"),
    ("PATINHO", "Patinho"),
]
# Itens que a fonte quebra em dois blocos, IN NATURA e PROCESSADO. O nome do
# item esta na linha de secao; o lado, no titulo do bloco.
DIVIDIDOS = [("ANEL DE LULA", "Anel de lula"),
             ("TENTACULO DE LULA", "Tentáculo de lula"),
             ("LOMBO DE ATUM", "Lombo de atum")]
# Blocos que existem na fonte e nao sao item de estoque comprado.
FORA = ["FILE PRODUCAO", "PROCESSAMENTO", "CAMARAO P (GERAL)"]

# Ordem das linhas = ordem de leitura dos blocos na fonte (linha a linha, da
# coluna B para a K). E a mesma ordem que a planilha de precos ja usava.
ORDEM = [
    "Salmão — libra 08/10", "Salmão — libra 10/12",
    "Salmão — libra 12/14", "Salmão — libra 14/16",
    "Salmão — raspa", "Salmão — apara", "Salmão — pele",
    "Camarão G s/rabo", "Camarão M s/rabo",
    "Camarão P s/rabo", "Camarão P c/rabo",
    "Kani", "Anchova", "Nori", "Cream cheese",
    "Arroz de sushi", "Anel de lula — in natura", "Anel de lula — processado",
    "Filé de tilápia", "Tentáculo de lula — in natura", "Tentáculo de lula — processado",
    "Patinho", "Lombo de atum — in natura", "Lombo de atum — processado",
]
POS = {n: i for i, n in enumerate(ORDEM)}
CASA_POS = {c: i for i, (c, _) in enumerate(CASAS)}
NOME_CASA = dict(CASAS)

# O nome das abas carrega o mes. Ficava fixo em 'agosto': em setembro o script
# nao achava aba nenhuma e pulava as duas caladamente.
def nomes_abas(mes):
    m = MESES[mes].lower()
    return "2. Entradas de %s" % m, "3. Transferências de %s" % m


def mes_anterior(mes, ano):
    return (12, ano - 1) if mes == 1 else (mes - 1, ano)


def nome_aba_saldo(mes, ano):
    """'1. Saldo inicial (31-07)' — o ultimo dia do mes anterior."""
    fim = datetime(ano, mes, 1) - timedelta(days=1)
    return "1. Saldo inicial (%s)" % fim.strftime("%d-%m")


# Unidade de compra do item que nao apareceu em nenhum dia do mes — sem isso a
# linha dele na aba de saldo sairia com a coluna 'Un.' vazia.
UNIDADE_PADRAO = {"Nori": "Pct. c/50Fls", "Cream cheese": "Bng c/1,5Kg",
                  "Arroz de sushi": "Pct. c/5Kg"}


def unidade_padrao(item):
    return UNIDADE_PADRAO.get(item, "Pxs" if "libra" in item else "Kg")


# Colunas de cada aba: (cabecalho, largura).
COLS_ENTRADAS = [("Data", 12), ("Casa", 9), ("Unidade", 22), ("Item", 34),
                 ("Un.", 16), ("Quantidade", 14), ("Preço unitário (R$)", 20),
                 ("Valor total", 18)]
COLS_TRANSF = [("Data", 12), ("Casa que enviou", 18), ("Unidade", 22),
               ("Item", 34), ("Un.", 16), ("Quantidade", 14),
               ("Valor unitário (R$)", 20), ("Valor total", 18)]


def cols_saldo(mes):
    anterior = MESES[mes_anterior(mes, 0)[0]].lower()
    return [("Casa", 9), ("Unidade", 22), ("Item", 34), ("Un.", 16),
            ("Saldo final de %s" % anterior, 20), ("Valor unitário (R$)", 20),
            ("Valor total", 18)]


# Formatos herdados da planilha, para a linha nova sair igual a que ja existe.
FMT_DATA = "mm-dd-yy"                    # builtin 14: o Excel mostra no padrao local
FMT_QTD = "#,##0.000"
FMT_MOEDA = r"\R\$\ #,##0.00"
LARANJA = "FFFFF4E6"                     # a coluna que o financeiro preenche
CHUMBO = "FF1F2430"                      # fundo do cabecalho
COR_TITULO = "FFF36F1D"
COR_SUBTITULO = "FF666666"


def sem_acento(s):
    t = unicodedata.normalize("NFKD", str(s or ""))
    return "".join(c for c in t if not unicodedata.combining(c)).upper().strip()


def numero(v):
    return float(v) if isinstance(v, (int, float)) else None


def campo_do_rotulo(rotulo):
    r = sem_acento(rotulo)
    for prefixo, campo in CAMPOS:
        if r.startswith(prefixo):
            return campo
    return None


def unidade_do_rotulo(rotulo):
    """'Início (Pct. c/50Fls):' -> 'Pct. c/50Fls' — a unidade de compra."""
    m = re.search(r"\(([^)]*)\)", str(rotulo or ""))
    return m.group(1).strip() if m else None


def nome_do_bloco(titulo, secao):
    """Nome do item, ou None quando o bloco nao e item de estoque."""
    t, sc = sem_acento(titulo), sem_acento(secao)
    for alvo in FORA:
        if t.startswith(alvo):
            return None
    for prefixo, nome in SIMPLES:
        if t.startswith(prefixo):
            return nome
    lado = "in natura" if t.startswith("IN NATURA") else (
        "processado" if t.startswith("PROCESSADO") else None)
    if lado:
        for prefixo, base in DIVIDIDOS:
            if sc.startswith(prefixo):
                return "%s — %s" % (base, lado)
        return None
    # bloco sem nome proprio (o titulo e so ' USO = 0,000'): vale a secao
    for prefixo, nome in SIMPLES:
        if sc.startswith(prefixo):
            return nome
    return None


# Mesma espera do importar_diario.py: o OneDrive segura o arquivo por alguns
# segundos enquanto sincroniza, e a fila anda de arquivo em arquivo.
ESPERA_LOCK = [0.5, 1, 2, 3, 5, 8]


def abrir(arq):
    """Abre a planilha mesmo travada pelo OneDrive ou aberta no Excel.

    A copia simples nao resolve sozinha: o open() do Python pede o arquivo sem
    compartilhamento e o Excel/OneDrive recusa. Abrir pelo FileStream do .NET
    com FileShare.ReadWrite le por cima do lock — e o unico caminho que
    funciona com a planilha aberta na tela do usuario.
    """
    import openpyxl
    for espera in [0] + ESPERA_LOCK:
        if espera:
            time.sleep(espera)
        try:
            return openpyxl.load_workbook(arq, data_only=True)
        except PermissionError:
            pass
        try:
            import clr  # noqa: F401
        except ImportError:
            pass
        dados = _ler_compartilhado(arq)
        if dados is not None:
            import io
            return openpyxl.load_workbook(io.BytesIO(dados), data_only=True)
    return None


def _ler_compartilhado(arq):
    """Le os bytes do arquivo mesmo com lock de escrita de outro processo."""
    import subprocess
    ps = (
        "$fs=New-Object System.IO.FileStream("
        "'%s',[System.IO.FileMode]::Open,[System.IO.FileAccess]::Read,"
        "([System.IO.FileShare]::ReadWrite -bor [System.IO.FileShare]::Delete));"
        "$ms=New-Object System.IO.MemoryStream;$fs.CopyTo($ms);$fs.Close();"
        "[Console]::OpenStandardOutput().Write($ms.ToArray(),0,$ms.Length)"
    ) % str(arq).replace("'", "''")
    try:
        r = subprocess.run(["powershell", "-NoProfile", "-Command", ps],
                           capture_output=True, timeout=120)
    except (OSError, subprocess.TimeoutExpired):
        return None
    return r.stdout if r.returncode == 0 and r.stdout else None


def pode_escrever(arq):
    """O Excel aberto segura o arquivo: da para LER por cima do lock, gravar nao.

    Sem esta conferencia o script lia as quatro planilhas-fonte (dois minutos)
    e so morria no fim, num PermissionError cru do shutil.copy do backup.
    """
    try:
        with open(arq, "r+b"):
            return True
    except (PermissionError, OSError):
        return False


def copiar_backup(origem, alvo):
    """Copia para os backups mesmo com o OneDrive segurando o arquivo."""
    try:
        shutil.copy(origem, alvo)
        return True
    except PermissionError:
        dados = _ler_compartilhado(origem)
        if dados is None:
            return False
        alvo.write_bytes(dados)
        return True


def ler_dia(ws):
    """{item: {campo: valor, 'un': unidade}} de uma aba de dia."""
    secao = {c: "" for c in COLUNAS}
    itens = {}

    for r in range(10, ws.max_row + 1):
        # Linha de secao: manda nas colunas dali para a direita, ate a proxima
        # coluna que tenha titulo proprio. E o que separa o bloco PROCESSADO do
        # anel de lula da secao vizinha (nori), que fica na mesma linha.
        titulos = {}
        for c in COLUNAS:
            v = ws.cell(r, c).value
            if not isinstance(v, str) or not v.strip():
                continue
            # 'Início (Kg):', 'Produto In Natura (Kg):', 'Rendimento (%):' sao
            # ROTULOS, nao titulo de secao — todos terminam em dois-pontos, e os
            # titulos de verdade ('KANI - SUSHIBAR') nunca terminam assim.
            if v.strip().endswith(":"):
                continue
            abaixo = ws.cell(r + 1, c).value
            eh_bloco = isinstance(abaixo, str) and campo_do_rotulo(abaixo)
            if not eh_bloco and "USO =" not in v and not campo_do_rotulo(v):
                titulos[c] = v
        if titulos:
            atual = None
            for c in COLUNAS:
                if c in titulos:
                    atual = titulos[c]
                if atual:
                    secao[c] = atual

        for c in COLUNAS:
            titulo = ws.cell(r, c).value
            if not isinstance(titulo, str) or not titulo.strip():
                continue
            # Sem esta guarda, 'Transferência (Kg):' passava por inicio de bloco
            # — a linha de baixo ('Perda') e rotulo valido —, caia no nome da
            # secao e relia os rotulos DESLOCADOS. Era assim que o Kani herdava
            # o 'Início (Pct. c/5Kg): 71' do arroz, tres linhas abaixo.
            if campo_do_rotulo(titulo):
                continue
            abaixo = ws.cell(r + 1, c).value
            if not (isinstance(abaixo, str) and campo_do_rotulo(abaixo)):
                continue
            nome = nome_do_bloco(titulo, secao[c])
            if not nome:
                continue
            reg = itens.setdefault(nome, {})
            for k in range(1, 7):
                rot = ws.cell(r + k, c).value
                if not isinstance(rot, str):
                    continue
                campo = campo_do_rotulo(rot)
                if not campo:
                    continue
                un = unidade_do_rotulo(rot)
                if un:
                    reg["un"] = un
                val = numero(ws.cell(r + k, c + 1).value)
                if val is not None:
                    reg[campo] = round(val, 4)
    return itens


def arquivo_da_casa(casa, mes):
    pasta = ORIGEM / casa / ("%02d - %s" % (mes, MESES[mes]))
    if not pasta.is_dir():
        return None
    achados = sorted(pasta.glob("*.xlsm"))
    return achados[0] if achados else None


def ler_fonte(mes, ano):
    """{(casa, 'YYYY-MM-DD', item): campos} com TODOS os itens da fonte."""
    fora, faltando = {}, []
    for casa, _nome in CASAS:
        arq = arquivo_da_casa(casa, mes)
        if not arq:
            faltando.append(casa)
            continue
        wb = abrir(arq)
        if wb is None:
            faltando.append("%s (travada)" % casa)
            continue
        for aba in wb.sheetnames:
            if not re.fullmatch(r"\d{2}", aba):
                continue
            try:
                iso = datetime(ano, mes, int(aba)).strftime("%Y-%m-%d")
            except ValueError:
                continue                  # dia 31 em mes de 30: aba sobrando
            for item, campos in ler_dia(wb[aba]).items():
                if item not in POS:
                    continue
                fora[(casa, iso, item)] = campos
    return fora, faltando


def linhas(fonte, campo):
    """As linhas da aba, na ordem: casa, data, ordem do item na fonte."""
    fora = []
    for (casa, iso, item), campos in fonte.items():
        qtd = campos.get(campo)
        if not qtd:
            continue
        fora.append((casa, iso, item, qtd, campos.get("un") or ""))
    fora.sort(key=lambda t: (CASA_POS[t[0]], t[1], POS[t[2]]))
    return fora


def saldo_final(mes, ano):
    """[(casa, item, qtd, un)] — o fechamento do mes, na ordem da aba 1.

    Vale o 'Final' do ultimo dia em que o item aparece, e nao o do dia 31: item
    que a casa parou de movimentar no meio do mes nao tem bloco no ultimo dia.
    Item que nao apareceu em dia nenhum entra com zero, como na planilha de
    agosto — a linha existe para o financeiro ver que o saldo e zero.

    Conferido contra a aba de saldo inicial de agosto, que foi montada por fora:
    recalculado a partir de julho, as 96 linhas (4 casas x 24 itens) batem valor
    por valor, na mesma ordem.
    """
    fonte, faltando = ler_fonte(mes, ano)
    ultimo = {}
    for (casa, iso, item), campos in fonte.items():
        if "final" not in campos:
            continue
        atual = ultimo.get((casa, item))
        if atual is None or iso > atual[0]:
            ultimo[(casa, item)] = (iso, campos["final"], campos.get("un"))
    fora = []
    for casa, _nome in CASAS:
        for item in sorted(ORDEM, key=sem_acento):   # a ordem que a aba 1 usa
            _iso, qtd, un = ultimo.get((casa, item), (None, 0, None))
            fora.append((casa, item, qtd, un or unidade_padrao(item)))
    return fora, faltando


def precos_digitados(ws):
    """{(casa, data, item): preco} do que ja esta preenchido na planilha."""
    fora = {}
    for row in ws.iter_rows(min_row=5):
        data, casa, item, preco = row[0].value, row[1].value, row[3].value, row[6].value
        if data is None or not hasattr(data, "strftime"):
            continue
        if preco is not None:
            fora[(casa, data.strftime("%Y-%m-%d"), item)] = preco
    return fora


def borda_fina():
    from openpyxl.styles import Border, Side
    fina = Side(style="thin")
    return Border(left=fina, right=fina, top=fina, bottom=fina)


def cabecalho(ws, titulo, subtitulo, colunas):
    """Linhas 1, 2 e 4 no mesmo padrao das abas que ja existiam."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    ws.cell(1, 1).value = titulo
    ws.cell(1, 1).font = Font(name="Calibri", sz=14, b=True, color=COR_TITULO)
    ws.row_dimensions[1].height = 18.75
    ws.cell(2, 1).value = subtitulo
    ws.cell(2, 1).font = Font(name="Calibri", sz=10, color=COR_SUBTITULO)
    for i, (nome, largura) in enumerate(colunas):
        cel = ws.cell(4, 1 + i)
        cel.value = nome
        cel.font = Font(name="Calibri", sz=10, b=True, color="FFFFFFFF")
        cel.fill = PatternFill("solid", fgColor=CHUMBO)
        cel.alignment = Alignment(horizontal="center", vertical="center",
                                  wrap_text=True)
        ws.column_dimensions[get_column_letter(1 + i)].width = largura


def escrever_saldo(ws, saldos, guardados):
    """Aba 1: uma linha por casa e item, com o saldo que virou o mes."""
    from openpyxl.styles import PatternFill
    laranja = PatternFill("solid", fgColor=LARANJA)
    borda = borda_fina()
    for i, (casa, item, qtd, un) in enumerate(saldos):
        r = 5 + i
        ws.cell(r, 1).value = casa
        ws.cell(r, 2).value = NOME_CASA[casa]
        ws.cell(r, 3).value = item
        ws.cell(r, 4).value = un
        saldo = ws.cell(r, 5)
        saldo.value = qtd
        saldo.number_format = FMT_QTD
        valor = ws.cell(r, 6)
        valor.value = guardados.get((casa, item))
        valor.number_format = FMT_MOEDA
        valor.fill = laranja
        total = ws.cell(r, 7)
        total.value = '=IF(F{0}="","",E{0}*F{0})'.format(r)
        total.number_format = FMT_MOEDA
        for c in range(1, 8):
            ws.cell(r, c).border = borda
    ultima = 4 + len(saldos)
    ws.auto_filter.ref = "A4:G%d" % max(ultima, 5)
    ws.freeze_panes = "A5"
    return ultima


def criar(destino, mes, ano, fonte, saldos):
    """Monta o arquivo do mes: saldo inicial, entradas e transferencias."""
    import openpyxl
    aba_ent, aba_tr = nomes_abas(mes)
    nome_mes = MESES[mes].lower()
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = nome_aba_saldo(mes, ano)
    cabecalho(ws, ws.title,
              "ETAPA 2 — Com quanto %s começa. Preencha o valor unitário de "
              "cada saldo na coluna laranja." % nome_mes, cols_saldo(mes))
    escrever_saldo(ws, saldos, {})

    ws = wb.create_sheet(aba_ent)
    cabecalho(ws, aba_ent,
              "ETAPA 1 — Cada chegada de produto em %s. Preencha o preço "
              "unitário real da compra na coluna laranja." % nome_mes,
              COLS_ENTRADAS)
    escrever_aba(ws, linhas(fonte, "entrada"), {})

    ws = wb.create_sheet(aba_tr)
    cabecalho(ws, aba_tr,
              "Saídas de uma casa para outra. O valor costuma ser o mesmo da "
              "entrada correspondente.", COLS_TRANSF)
    escrever_aba(ws, linhas(fonte, "transferencia"), {})

    wb.save(destino)


def escrever_aba(ws, dados, guardados):
    """Reescreve a aba inteira, devolvendo os precos que ja estavam digitados."""
    from openpyxl.styles import PatternFill
    laranja = PatternFill("solid", fgColor=LARANJA)
    borda = borda_fina()

    # ATENCAO: ws.cell(r, c, valor) do openpyxl IGNORA a atribuicao quando o
    # valor e None ("if value is not None"). Como as linhas mudam de lugar a
    # cada rodada (os dias novos entram no meio, por casa e data), a celula de
    # preco de uma linha sem preco ficava com o valor do item que ANTES ocupava
    # aquela linha — 74 precos trocados de item na primeira tentativa. Por isso
    # aqui todo valor e escrito por '.value =', que aceita None.
    for i, (casa, iso, item, qtd, un) in enumerate(dados):
        r = 5 + i
        data = ws.cell(r, 1)
        data.value = datetime.strptime(iso, "%Y-%m-%d")
        data.number_format = FMT_DATA
        ws.cell(r, 2).value = casa
        ws.cell(r, 3).value = NOME_CASA[casa]
        ws.cell(r, 4).value = item
        ws.cell(r, 5).value = un
        quantidade = ws.cell(r, 6)
        quantidade.value = qtd
        quantidade.number_format = FMT_QTD
        preco = ws.cell(r, 7)
        preco.value = guardados.get((casa, iso, item))
        preco.number_format = FMT_MOEDA
        preco.fill = laranja
        total = ws.cell(r, 8)
        total.value = '=IF(G{0}="","",F{0}*G{0})'.format(r)
        total.number_format = FMT_MOEDA
        # A linha que nasce depois do fim da aba antiga vinha sem borda: a
        # grade parava no meio da aba e o resto ficava solto na tela.
        for c in range(1, 9):
            ws.cell(r, c).border = borda

    ultima = 4 + len(dados)
    if ws.max_row > ultima:                    # limpa a sobra de uma rodada maior
        ws.delete_rows(ultima + 1, ws.max_row - ultima)
    ws.auto_filter.ref = "A4:H%d" % max(ultima, 5)
    ws.freeze_panes = "A5"                     # estava em A190: congelava meia aba
    return ultima


def montar(destino, mes, ano, fonte, gravar):
    """Cria a planilha do mes. O saldo inicial vem do fechamento do mes anterior."""
    m_ant, a_ant = mes_anterior(mes, ano)
    print("\nplanilha nova. saldo inicial = fechamento de %s %d (lendo a fonte)"
          % (MESES[m_ant], a_ant))
    saldos, faltando = saldo_final(m_ant, a_ant)
    if faltando:
        print("[aviso] no mes anterior nao li: %s" % ", ".join(faltando))
        print("        essas casas entram com saldo zero — confira antes de"
              "\n        mandar a planilha para preenchimento.")
    aba_ent, aba_tr = nomes_abas(mes)
    entradas, transf = linhas(fonte, "entrada"), linhas(fonte, "transferencia")
    com_saldo = sum(1 for _c, _i, q, _u in saldos if q)

    print()
    print("%-34s %3d linhas (%d com saldo, %d zerados)"
          % (nome_aba_saldo(mes, ano), len(saldos), com_saldo,
             len(saldos) - com_saldo))
    print("%-34s %3d linhas" % (aba_ent, len(entradas)))
    print("%-34s %3d linhas" % (aba_tr, len(transf)))
    print("\ncoluna laranja toda em branco: %d valores a preencher"
          % (len(saldos) + len(entradas) + len(transf)))

    if not gravar:
        print("\n(nada foi gravado — rode com --gravar)")
        return 0
    criar(destino, mes, ano, fonte, saldos)
    print("\ncriada: %s" % destino.name)
    return 0


def main():
    args = sys.argv[1:]
    gravar = "--gravar" in args
    criar_nova = "--criar" in args
    hoje = datetime.now()
    mes, ano = hoje.month, hoje.year
    if "--mes" in args:
        mes = int(args[args.index("--mes") + 1])

    destino = (RAIZ / ".." / ("Cadastro de Precos - %s %d.xlsx" % (MESES[mes], ano))).resolve()
    print("fonte  : %s" % ORIGEM)
    print("destino: %s" % destino)
    if not destino.exists() and not criar_nova:
        print("\n[erro] a planilha nao existe. Para montar o mes do zero:"
              "\n       python gerar_cadastro_precos.py --mes %02d --criar --gravar"
              % mes)
        return 1

    if gravar and destino.exists() and not pode_escrever(destino):
        print("\n[erro] a planilha esta aberta — o Excel (ou o OneDrive) esta"
              "\n       segurando o arquivo e nao da para gravar nele."
              "\n       Feche ela e rode de novo.")
        return 1

    fonte, faltando = ler_fonte(mes, ano)
    if faltando:
        print("\n[aviso] nao li: %s" % ", ".join(faltando))
    if not fonte:
        print("\n[erro] nenhuma planilha-fonte lida.")
        return 1

    if not destino.exists():
        return montar(destino, mes, ano, fonte, gravar)
    if criar_nova:
        print("\n[erro] a planilha ja existe — o --criar refaria o saldo inicial"
              "\n       por cima do que ja foi digitado. Rode sem --criar para"
              "\n       atualizar as abas de entradas e transferencias.")
        return 1

    import openpyxl
    dados_planilha = _ler_compartilhado(destino)
    if dados_planilha is None:
        print("\n[erro] nao consegui ler a planilha de precos.")
        return 1
    import io
    wb = openpyxl.load_workbook(io.BytesIO(dados_planilha), data_only=False)

    relato = []
    aba_entradas, aba_transf = nomes_abas(mes)
    for aba, campo in [(aba_entradas, "entrada"), (aba_transf, "transferencia")]:
        if aba not in wb.sheetnames:
            print("\n[aviso] aba '%s' nao existe — pulei." % aba)
            continue
        ws = wb[aba]
        antes = precos_digitados(ws)
        dados = linhas(fonte, campo)
        novas = [d for d in dados if (d[0], d[1], d[2]) not in antes]
        relato.append((aba, len(dados), len(novas), novas))
        if gravar:
            escrever_aba(ws, dados, antes)

    print()
    for aba, total, n_novas, novas in relato:
        print("%-32s %3d linhas | %3d novas (sem preco)" % (aba, total, n_novas))
        for casa, iso, item, qtd, un in novas:
            print("    %s  %-5s %-32s %10.3f %s" % (iso[8:] + "/" + iso[5:7],
                                                    casa, item, qtd, un))

    if not gravar:
        print("\n(nada foi gravado — rode com --gravar)")
        return 0

    BACKUPS.mkdir(exist_ok=True)
    carimbo = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    if not copiar_backup(destino, BACKUPS / ("cadastro_precos_%s.xlsx" % carimbo)):
        print("\n[erro] nao consegui fazer o backup da planilha — nao gravei nada."
              "\n       Feche a planilha e rode de novo.")
        return 1
    wb.save(destino)
    print("\ngravado. backup em backups/cadastro_precos_%s.xlsx" % carimbo)
    return 0


if __name__ == "__main__":
    sys.exit(main())
