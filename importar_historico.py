# -*- coding: utf-8 -*-
"""
Importa o historico das 32 planilhas do CDE (4 casas x 8 meses) para o CDE Web.

Le as abas CONSUMO - *, que ja trazem a serie diaria consolidada.
Por padrao so CONFERE e mostra o relatorio; passe --gravar para escrever de fato.

    python importar_historico.py             confere e mostra o relatorio
    python importar_historico.py --gravar    grava em dados/lancamentos.json
"""
import json, re, sys, unicodedata
from datetime import datetime
from pathlib import Path

RAIZ = Path(__file__).parent
ORIGEM = (RAIZ / ".." / ".." / ".." / "Documentos 2026" / "CDE - Controle Diário de Estoque").resolve()
DESTINO = RAIZ / "dados" / "lancamentos.json"

CASAS = ["SL", "DC", "DLU", "DLCN"]

# fatores de conversao para peixe-equivalente 08/10 (confirmados pela area)
FATOR_LIBRA = {"08/10": 1.00, "10/12": 1.25, "12/14": 1.48, "14/16": 1.67}

# aba de consumo -> chave do insumo no sistema novo
MAPA = {
    "salmao": "salmao_equivalente",
    "carmao g": "camarao_g", "camarao g": "camarao_g",
    "carmao m": "camarao_m", "camarao m": "camarao_m",
    "carmao p - s rabo": "camarao_p", "camarao p - s rabo": "camarao_p",
    "kani": "kani", "anchova": "anchova", "nori": "nori",
    "cream cheese": "cream_cheese", "arroz de sushi": "arroz",
    "file de tilapia": "file_tilapia", "patinho": "patinho",
    "lula": "lula", "polvo": "polvo", "atum": "atum",
}
# abas que existem mas nao entram (o usuario pediu so o camarao sem rabo)
IGNORAR = {"carmao p - c rabo", "camarao p - c rabo", "camarao p (geral)", "carmao p (geral)"}


def sem_acento(s):
    s = unicodedata.normalize("NFKD", str(s))
    return "".join(c for c in s if not unicodedata.combining(c)).lower().strip()


def chave_da_aba(nome):
    if not nome.upper().startswith("CONSUMO - "):
        return None
    resto = sem_acento(nome[10:])
    if resto in IGNORAR:
        return None
    return MAPA.get(resto)


def achar_colunas(ws):
    """Localiza as colunas pelo texto do cabecalho (linha 4), que varia entre abas."""
    col = {}
    for c in range(1, ws.max_column + 1):
        h = sem_acento(ws.cell(row=4, column=c).value or "")
        if not h:
            continue
        if h.startswith("data"):                      col["data"] = c
        elif h.startswith("faturamento"):             col["faturamento"] = c
        elif "estoque inicial" in h:                  col["inicial"] = c
        elif "estoque final" in h:                    col["final"] = c
        elif "entrada" in h or "saida" in h:          col["entrada"] = c
        elif h.startswith("disperdicio") or h.startswith("desperdicio"): col["desperdicio"] = c
        elif h.startswith("uso diario"):              col["uso"] = c
        elif "coeficiente" in h:                      col["coef"] = c
    return col


def numero(v):
    return float(v) if isinstance(v, (int, float)) else None


def abertas_no_excel():
    """Planilhas que estao abertas no Excel neste momento.

    Ao abrir um arquivo, o Excel cria ao lado dele um arquivo de posse com o
    prefixo "~$". Como cada pasta de mes tem exatamente um .xlsm, a presenca
    desse arquivo na pasta identifica a planilha aberta sem ambiguidade — o
    nome do arquivo de posse nem sempre repete o nome completo do original.
    """
    abertas = []
    for trava in ORIGEM.glob("*/*/~$*"):
        # o proprio arquivo de posse termina em .xlsm (e "~$" + o nome original),
        # entao ele entra no glob e precisa ser descartado
        abertas.extend(x for x in trava.parent.glob("*.xlsm")
                       if not x.name.startswith("~$"))
    return sorted(set(abertas))


def importar(arquivos=None):
    """Le as planilhas e devolve os lancamentos.

    Sem argumento le todas — e a importacao inicial do historico. O botao
    Sincronizar passa so as planilhas abertas, para nao reprocessar meses
    que ninguem esta mexendo.
    """
    import openpyxl
    if not ORIGEM.exists():
        raise SystemExit("Pasta de origem nao encontrada:\n  %s" % ORIGEM)

    dados, relato = {}, []
    divergencias, faixa_1416 = [], []
    if arquivos is None:
        arquivos = sorted(ORIGEM.glob("*/*/*.xlsm"))
        if not arquivos:
            raise SystemExit("Nenhuma planilha .xlsm encontrada em %s" % ORIGEM)
    arquivos = [Path(a) for a in arquivos]

    for arq in arquivos:
        casa = arq.parts[-3]
        if casa not in CASAS:
            continue
        try:
            wb = openpyxl.load_workbook(arq, data_only=True)
        except Exception as e:
            relato.append((casa, arq.parts[-2], 0, 0, "ERRO: %s" % e))
            continue

        dias_arq = linhas_arq = 0
        for aba in wb.sheetnames:
            chave = chave_da_aba(aba)
            if not chave:
                continue
            ws = wb[aba]
            col = achar_colunas(ws)
            if "data" not in col or "final" not in col:
                continue
            for r in range(5, ws.max_row + 1):
                d = ws.cell(row=r, column=col["data"]).value
                if not isinstance(d, datetime):
                    continue
                iso = d.strftime("%Y-%m-%d")
                final = numero(ws.cell(row=r, column=col["final"]).value)
                entrada = numero(ws.cell(row=r, column=col.get("entrada", 0)).value) if col.get("entrada") else None
                desp = numero(ws.cell(row=r, column=col.get("desperdicio", 0)).value) if col.get("desperdicio") else None
                fat = numero(ws.cell(row=r, column=col.get("faturamento", 0)).value) if col.get("faturamento") else None
                inicial = numero(ws.cell(row=r, column=col.get("inicial", 0)).value) if col.get("inicial") else None
                uso_planilha = numero(ws.cell(row=r, column=col.get("uso", 0)).value) if col.get("uso") else None

                # Dia ainda nao preenchido: a planilha deixa o estoque final em zero e
                # a formula interpreta isso como "consumiu tudo", gerando uso falso.
                # Mesma regra que o CDE ja usa na coluna Valido: so vale com faturamento.
                if not fat:
                    continue
                # dia sem nenhum movimento: nao importa (evita encher de zeros)
                if not any(v not in (None, 0) for v in (final, entrada, desp, uso_planilha)):
                    continue

                dia = dados.setdefault(iso, {})
                lan = dia.setdefault(casa, {"faturamento": fat, "insumos": {}})
                if fat not in (None, 0):
                    lan["faturamento"] = fat
                # O uso vem PRONTO da planilha, nao recalculado: lula, salmao, polvo e
                # atum passam por processamento e tem um rendimento embutido que a conta
                # "inicial + entrada - final" nao reproduz. Importamos o numero que a
                # area ja usa e confere ha meses.
                campos = {}
                if final is not None:   campos["final"] = round(final, 4)
                # A planilha tem uma coluna so para os dois sentidos: chegada
                # entra positiva, transferencia para outra casa entra negativa.
                # Guardar tudo como "entrada" dava chegada de -8 kg na tela e
                # deixava a coluna de saida sempre vazia.
                if entrada is not None:
                    if entrada < 0: campos["transferencia"] = round(-entrada, 4)
                    else:           campos["entrada"] = round(entrada, 4)
                # Na aba do salmao a coluna DESPERDICIO repete o estoque final
                # (formula errada na origem). Importado assim, o salmao aparecia
                # perdendo centenas de quilos por mes. Espelho nao e perda.
                if desp is not None and not (final is not None and desp
                                             and abs(desp - final) < 1e-9):
                    campos["desperdicio"] = round(desp, 4)
                if uso_planilha is not None: campos["uso"] = round(uso_planilha, 4)
                if campos:
                    lan["insumos"][chave] = campos
                    linhas_arq += 1

                # conferencia: o uso que eu recalculo bate com o da planilha?
                if None not in (inicial, final, uso_planilha):
                    meu = inicial + (entrada or 0) - final
                    if abs(meu - uso_planilha) > 0.01:
                        divergencias.append((casa, iso, chave, uso_planilha, round(meu, 3)))
            dias_arq = max(dias_arq, 0)

        # a faixa 14/16 chegou a ser usada? (4a coluna do bloco do salmao nas abas de dia)
        for aba in wb.sheetnames:
            if re.fullmatch(r"\d{2}", aba):
                ws = wb[aba]
                vals = [ws.cell(row=r, column=12).value for r in (16, 17, 18, 19)]
                if any(isinstance(x, (int, float)) and x != 0 for x in vals):
                    faixa_1416.append("%s %s dia %s" % (casa, arq.parts[-2], aba))
        wb.close()
        relato.append((casa, arq.parts[-2], linhas_arq, len(wb.sheetnames), ""))

    return dados, relato, divergencias, faixa_1416


if __name__ == "__main__":
    gravar = "--gravar" in sys.argv
    dados, relato, divergencias, faixa_1416 = importar()

    print("=" * 68)
    print("  IMPORTACAO DO HISTORICO — %s" % ("GRAVANDO" if gravar else "CONFERENCIA (nada foi gravado)"))
    print("=" * 68)
    print("\n  %-6s %-14s %8s" % ("CASA", "MES", "LANCAM."))
    for casa, mes, linhas, _abas, erro in relato:
        print("  %-6s %-14s %8s %s" % (casa, mes, linhas, erro))

    datas = sorted(dados)
    print("\n  dias com dados: %d   de %s a %s" % (len(datas), datas[0] if datas else "-", datas[-1] if datas else "-"))
    total = sum(len(c.get("insumos", {})) for d in dados.values() for c in d.values() if isinstance(c, dict))
    print("  lancamentos de insumo: %d" % total)

    print("\n  CONFERENCIA — uso recalculado x uso da planilha")
    if divergencias:
        print("  %d divergencia(s) acima de 0,01. Primeiras 12:" % len(divergencias))
        for casa, iso, chave, planilha, meu in divergencias[:12]:
            print("     %-5s %s %-20s planilha=%-10s recalculado=%s" % (casa, iso, chave, planilha, meu))
    else:
        print("  nenhuma divergencia — todos os usos batem com a planilha")

    print("\n  FAIXA 14/16 (a que estava com fator 1 na formula)")
    print("  %s" % ("nunca usada — o erro ficou latente, nao contaminou numero nenhum"
                    if not faixa_1416 else "USADA em: " + ", ".join(faixa_1416[:10])))

    if gravar:
        DESTINO.parent.mkdir(exist_ok=True)
        if DESTINO.exists():
            DESTINO.replace(DESTINO.with_suffix(".json.antes-da-importacao"))
        with open(DESTINO, "w", encoding="utf-8") as f:
            json.dump(dados, f, ensure_ascii=False, indent=1)
        print("\n  gravado em %s" % DESTINO)
    else:
        print("\n  Nada foi gravado. Rode com --gravar para escrever.")
