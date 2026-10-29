# -*- coding: utf-8 -*-
"""
Base do CDE por casa: um .xlsx por casa e por ano, gravado pelo sistema.

Por que existe: as planilhas mensais (.xlsm) tem macros sensiveis, e a partir
de novembro/2026 o controle diario deixa de ser digitado nelas. O lancamento
passa a ser feito so na tela de Lancamento Diario, e cada dia salvo vai para
"dados/consolidado/CDE <ano> - <CASA>.xlsx", em formato de tabela (uma linha
por dia x insumo). Ninguem digita nesse arquivo: ele guarda o que o sistema
gravou.

Quem escreve e o servidor (vigiar_e_exportar, numa thread): quando o
lancamentos.json muda, os arquivos sao refeitos inteiros a partir dele. Refazer
inteiro, e nao acrescentar linha, garante que uma correcao num dia antigo
chegue ao arquivo tambem. O historico de janeiro a outubro veio das planilhas
pela sincronizacao e foi conferido linha a linha contra elas.

Os responsaveis do dia e as anotacoes existem so nas planilhas antigas: sao
lidos uma vez (--historico) e ficam em dados/consolidado/dias_planilha.json,
de onde a aba DIAS os traz de volta a cada exportacao.

Uso:
    python consolidar_cde.py               refaz os arquivos a partir do sistema
    python consolidar_cde.py --historico   le responsaveis e anotacoes das .xlsm
"""
import json
import re
import sys
import threading
import time
from datetime import datetime, timedelta
from pathlib import Path

RAIZ = Path(__file__).parent
DESTINO = RAIZ / "dados" / "consolidado"
ARQ_DIAS = DESTINO / "dias_planilha.json"

CABECALHO_MOV = ["data", "insumo", "item", "unidade", "inicial", "chegada",
                 "transferencia", "perda", "final", "uso", "origem"]
CABECALHO_DIA = ["data", "faturamento", "resp_unidade", "resp_sushibar", "resp_cozinha",
                 "anotacoes"]
CABECALHO_PROC = ["data", "insumo", "in_natura", "processado", "perda", "rendimento"]


# ------------------------------------------------------------------ exportacao
def linhas_da_casa(dados, ck, ano, dias_planilha):
    """(movimento, dias, processamento) de uma casa num ano, do lancamentos.json."""
    import server
    catalogo = [i for i in server.INSUMOS if not i.get("derivado")]
    mov, dias, procs = [], [], []
    final_anterior = {}
    for iso in sorted(k for k in dados if k[:4] == str(ano)):
        lan = (dados.get(iso) or {}).get(ck)
        if not isinstance(lan, dict):
            continue
        ins = lan.get("insumos") or {}
        data = datetime.strptime(iso, "%Y-%m-%d")
        teve = False
        for item in catalogo:
            c = ins.get(item["chave"])
            if not isinstance(c, dict) or not any(
                    c.get(k) is not None for k in server.CAMPOS_LANC):
                continue
            # o inicial e o final do dia anterior quando a tela nao gravou um
            ini = c.get("inicial")
            if ini is None:
                ini = final_anterior.get(item["chave"])
            fim = c.get("final")
            uso = c.get("uso")
            if uso is None and fim is not None and ini is not None \
                    and not item.get("so_contagem"):
                uso = round(server.uso_do_dia(ini, c.get("entrada"), c.get("transferencia"),
                                              c.get("desperdicio"), fim), 4)
            origem = ("planilha" if c.get("_fonte") == "planilha"
                      else "editado no sistema" if c.get("_editado") else "sistema")
            mov.append([data, item["chave"], item["nome"], item["un"], ini, c.get("entrada"),
                        c.get("transferencia"), c.get("desperdicio"), fim, uso, origem])
            if fim is not None:
                final_anterior[item["chave"]] = fim
            teve = True
        for pk, p in (lan.get("processamento") or {}).items():
            cru, feito = p.get("in_natura"), p.get("processado")
            if not cru:
                continue
            procs.append([data, pk, cru, feito,
                          round(cru - (feito or 0), 4),
                          round((feito or 0) / cru, 4)])
        if teve or lan.get("faturamento"):
            extra = (dias_planilha.get(ck) or {}).get(iso) or {}
            dias.append([data, lan.get("faturamento")]
                        + [extra.get(k) for k in CABECALHO_DIA[2:]])
    return mov, dias, procs


def gravar_xlsx(casa, ano, mov, dias, procs, arquivo):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    abas = [("MOVIMENTO", CABECALHO_MOV, mov), ("DIAS", CABECALHO_DIA, dias),
            ("PROCESSAMENTO", CABECALHO_PROC, procs)]
    for i, (nome, cab, linhas) in enumerate(abas):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = nome
        ws.append(cab)
        for linha in linhas:
            ws.append(linha)
        for cel in ws[1]:
            cel.font = Font(bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="1F3A5F")
        for cel in ws["A"][1:]:
            cel.number_format = "DD/MM/YYYY"
        for j, titulo in enumerate(cab):
            letra = ws.cell(row=1, column=j + 1).column_letter
            ws.column_dimensions[letra].width = 30 if titulo in ("item", "anotacoes") else 14
        if linhas:
            ref = "A1:%s%d" % (ws.cell(row=1, column=len(cab)).column_letter, len(linhas) + 1)
            tab = Table(displayName="%s_%s" % (nome, casa), ref=ref)
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
            ws.add_table(tab)
        ws.freeze_panes = "A2"
    sobre = wb.create_sheet("SOBRE")
    for linha in [["Casa", casa], ["Ano", ano],
                  ["Gravado em", datetime.now().strftime("%d/%m/%Y %H:%M:%S")],
                  ["Fonte", "Tela de Lancamento Diario do CDE Web (dados/lancamentos.json). "
                            "Nao digitar aqui: o arquivo e refeito a cada gravacao."],
                  ["Regra", "inicial = final do dia anterior, salvo correcao na conferencia; "
                            "uso = inicial + chegada - transferencia - perda - final "
                            "(perda nao e consumo). "
                            "origem: planilha (historico das .xlsm), sistema ou editado no sistema."]]:
        sobre.append(linha)
    sobre.column_dimensions["A"].width = 14
    sobre.column_dimensions["B"].width = 110
    DESTINO.mkdir(parents=True, exist_ok=True)
    tmp = arquivo.with_name(arquivo.stem + ".tmp.xlsx")
    wb.save(tmp)
    import server
    server.trocar(tmp, arquivo)            # mesma troca atomica, com espera do OneDrive


def exportar(dados=None):
    """Refaz os arquivos de todas as casas, em cada ano que tem lancamento."""
    import server
    dados = server.carregar() if dados is None else dados
    dias_planilha = {}
    if ARQ_DIAS.exists():
        with open(ARQ_DIAS, encoding="utf-8") as f:
            dias_planilha = json.load(f)
    feitos = []
    for ano in sorted({k[:4] for k in dados if re.fullmatch(r"\d{4}-\d\d-\d\d", k)}):
        for casa in server.CASAS:
            ck = casa["chave"]
            mov, dias, procs = linhas_da_casa(dados, ck, ano, dias_planilha)
            if not mov:
                continue
            arquivo = DESTINO / ("CDE %s - %s.xlsx" % (ano, ck))
            gravar_xlsx(ck, ano, mov, dias, procs, arquivo)
            feitos.append((arquivo.name, len(dias), len(mov)))
    return feitos


def vigiar_e_exportar(intervalo=5, calma=3):
    """Thread do servidor: refaz a base das casas quando o lancamentos.json muda.

    Vigiar o arquivo, em vez de chamar a exportacao em cada rota de gravacao,
    cobre todo mundo que escreve nele (Lancamento, Movimentacoes, precos,
    sincronizacao) sem esquecer nenhum. A exportacao espera o arquivo ficar
    `calma` segundos sem mudar: quem digita varios precos em sequencia gera uma
    exportacao so, e ela nao atrasa a resposta da tela.
    """
    import server

    def laco():
        exportado = None
        while True:
            time.sleep(intervalo)
            try:
                mt = server.ARQ.stat().st_mtime
                if mt == exportado or time.time() - mt < calma:
                    continue
                exportar()
                exportado = mt
            except Exception as e:          # nunca derruba o servidor
                print("  [base das casas] falha ao gravar: %s" % e, flush=True)

    t = threading.Thread(target=laco, name="base-das-casas", daemon=True)
    t.start()
    return t


# --------------------------------------------------- historico das planilhas
def depois_dos_dois_pontos(v):
    t = str(v or "")
    return t.split(":", 1)[1].strip() if ":" in t else ""


def ler_historico_dias():
    """Responsaveis e anotacoes de cada dia, das abas diarias das .xlsm."""
    from importar_diario import ORIGEM, CASAS, abrir, data_do_dia, numero, sem_acento
    fora = {}
    for casa in CASAS:
        for arq in sorted((ORIGEM / casa).glob("*/*.xlsm")):
            wb = abrir(arq)
            if wb is None:
                print("  travada: %s %s" % (casa, arq.parts[-2]))
                continue
            for aba in wb.sheetnames:
                if not re.fullmatch(r"\d{1,2}", aba.strip()):
                    continue
                ws = wb[aba]
                iso, _ = data_do_dia(arq, aba, ws.cell(row=7, column=12).value)
                if not iso or not numero(ws.cell(row=9, column=12).value):
                    continue
                notas = []
                for r in range(1, ws.max_row + 1):
                    if sem_acento(ws.cell(row=r, column=2).value).startswith(
                            "ANOTACOES SOBRE O CDE"):
                        notas = [str(ws.cell(row=rr, column=c).value).strip()
                                 for rr in range(r + 1, ws.max_row + 1) for c in range(2, 13)
                                 if isinstance(ws.cell(row=rr, column=c).value, str)
                                 and ws.cell(row=rr, column=c).value.strip()]
                        break
                fora.setdefault(casa, {})[iso] = {
                    "resp_unidade": depois_dos_dois_pontos(ws.cell(row=7, column=5).value),
                    "resp_sushibar": depois_dos_dois_pontos(ws.cell(row=9, column=5).value),
                    "resp_cozinha": depois_dos_dois_pontos(ws.cell(row=11, column=5).value),
                    "anotacoes": " | ".join(notas) or None,
                }
            wb.close()
            print("  %s %s" % (casa, arq.parts[-2]), flush=True)
    DESTINO.mkdir(parents=True, exist_ok=True)
    with open(ARQ_DIAS, "w", encoding="utf-8") as f:
        json.dump(fora, f, ensure_ascii=False, indent=1)
    return fora


def main():
    if "--historico" in sys.argv:
        h = ler_historico_dias()
        print("historico: %s" % ", ".join("%s %d dias" % (c, len(d)) for c, d in h.items()))
    for nome, nd, nm in exportar():
        print("%s: %d dias, %d linhas" % (nome, nd, nm))


if __name__ == "__main__":
    main()
