# -*- coding: utf-8 -*-
"""
Preenche a aba "1. Saldo inicial" com o custo medio do mes anterior.

Por que existe: o estoque que virou o mes foi comprado no mes que passou, e o
preco daquelas compras ja esta digitado na planilha anterior — linha por linha,
com data, casa, item e quantidade. Nao ha motivo para redigitar 96 valores a
mao quando eles saem de uma media.

Qual media: MEDIA PONDERADA PELA QUANTIDADE, por casa e por item.

    valor unitario = soma(qtd x preco) / soma(qtd)

Ponderada, e nao media simples dos precos: uma chegada de 200 kg e uma de 5 kg
nao pesam igual no custo do que esta na camara. E por casa, porque as quatro
compram de fornecedores diferentes e o preco varia entre elas (veja
dados/custos.json, onde o salmao ja tem um valor por casa).

Quando a casa nao comprou o item no mes, cai na media das quatro. Quando
ninguem comprou, a linha fica em branco — e o script confere que o saldo dela e
zero. Saldo zero sem preco nao atrapalha (zero vezes qualquer coisa e zero); se
aparecer saldo com quantidade e sem preco, ele avisa em vez de inventar numero.

So preenche celula VAZIA. Valor ja digitado nunca e sobrescrito.

Uso:
    python preencher_saldo_inicial.py --mes 09            mostra o que faria
    python preencher_saldo_inicial.py --mes 09 --gravar   preenche
"""
import sys
from datetime import datetime
from pathlib import Path

import gerar_cadastro_precos as G

RAIZ = Path(__file__).parent


def medias_do_mes(arq, aba):
    """({(casa, item): (media, qtd, unidades)}, {item: (media, qtd)})."""
    import io
    import openpyxl
    dados = G._ler_compartilhado(arq)
    if dados is None:
        return None, None
    ws = openpyxl.load_workbook(io.BytesIO(dados), data_only=True)[aba]

    acc, geral = {}, {}
    for r in range(5, ws.max_row + 1):
        item = ws.cell(r, 4).value
        if not item:
            continue
        casa, un = ws.cell(r, 2).value, ws.cell(r, 5).value
        qtd, preco = ws.cell(r, 6).value, ws.cell(r, 7).value
        # Preco em formula sem valor em cache chega aqui como None — o
        # importar_precos_planilha avisa sobre isso; aqui a linha so nao pesa.
        if not isinstance(preco, (int, float)) or not isinstance(qtd, (int, float)):
            continue
        if qtd <= 0:
            continue
        a = acc.setdefault((casa, item), [0.0, 0.0, set()])
        a[0] += qtd * preco
        a[1] += qtd
        a[2].add(un)
        g = geral.setdefault(item, [0.0, 0.0])
        g[0] += qtd * preco
        g[1] += qtd
    return ({k: (v[0] / v[1], v[1], v[2]) for k, v in acc.items() if v[1] > 0},
            {k: (v[0] / v[1], v[1]) for k, v in geral.items() if v[1] > 0})


def main():
    args = sys.argv[1:]
    gravar = "--gravar" in args
    if "--mes" not in args:
        print("informe o mes:  python preencher_saldo_inicial.py --mes 09")
        return 1
    mes = int(args[args.index("--mes") + 1])
    ano = datetime.now().year
    if "--ano" in args:
        ano = int(args[args.index("--ano") + 1])
    m_ant, a_ant = G.mes_anterior(mes, ano)

    destino = (RAIZ / ".." / ("Cadastro de Precos - %s %d.xlsx" % (G.MESES[mes], ano))).resolve()
    fonte = (RAIZ / ".." / ("Cadastro de Precos - %s %d.xlsx" % (G.MESES[m_ant], a_ant))).resolve()
    print("preencher: %s" % destino.name)
    print("com media: %s" % fonte.name)
    for arq in (destino, fonte):
        if not arq.exists():
            print("\n[erro] nao encontrei: %s" % arq.name)
            return 1
    if gravar and not G.pode_escrever(destino):
        print("\n[erro] a planilha esta aberta — feche ela e rode de novo.")
        return 1

    aba_ent, _ = G.nomes_abas(m_ant)
    acc, geral = medias_do_mes(fonte, aba_ent)
    if acc is None:
        print("\n[erro] nao consegui ler a planilha de %s." % G.MESES[m_ant])
        return 1
    print("\nmedias de %s: %d combinacoes casa+item, %d itens no geral"
          % (G.MESES[m_ant], len(acc), len(geral)))

    import io
    import openpyxl
    bytes_destino = G._ler_compartilhado(destino)
    if bytes_destino is None:
        print("\n[erro] nao consegui ler a planilha de destino.")
        return 1
    wb = openpyxl.load_workbook(io.BytesIO(bytes_destino), data_only=False)
    aba = G.nome_aba_saldo(mes, ano)
    if aba not in wb.sheetnames:
        print("\n[erro] aba '%s' nao existe na planilha." % aba)
        return 1
    ws = wb[aba]

    preencher, ja_tinha, avisos, zerados = [], [], [], []
    for r in range(5, ws.max_row + 1):
        casa, item = ws.cell(r, 1).value, ws.cell(r, 3).value
        if not item:
            continue
        un, saldo = ws.cell(r, 4).value, ws.cell(r, 6 - 1).value
        if ws.cell(r, 6).value not in (None, ""):
            ja_tinha.append((casa, item))
            continue
        propria = acc.get((casa, item))
        if propria:
            media, qtd, uns = propria
            de = "%s (%s un)" % (casa, round(qtd))
            # A media so vale se o preco de agosto e da MESMA unidade do saldo:
            # preco por bisnaga em cima de saldo em quilo seria dinheiro errado.
            if un not in uns:
                avisos.append((casa, item, "unidade do saldo e %r, a compra veio em %s"
                               % (un, sorted(uns))))
                continue
        elif item in geral:
            media, qtd = geral[item]
            de = "media das 4 casas"
        else:
            if saldo:
                avisos.append((casa, item, "tem saldo %.3f e ninguem comprou em %s"
                               % (saldo, G.MESES[m_ant])))
            else:
                zerados.append((casa, item))
            continue
        preencher.append((r, casa, item, round(media, 4), de))

    print()
    print("a preencher                        : %d" % len(preencher))
    print("ja tinham valor (nao mexo)         : %d" % len(ja_tinha))
    print("sem compra no mes anterior, saldo 0: %d" % len(zerados))
    print("avisos                             : %d" % len(avisos))
    for casa, item, motivo in avisos:
        print("   [aviso] %-5s %-32s %s" % (casa, item, motivo))

    por_casa = {}
    for _r, casa, _i, _m, de in preencher:
        por_casa.setdefault(casa, [0, 0])
        por_casa[casa][0 if de != "media das 4 casas" else 1] += 1
    print()
    for casa in por_casa:
        print("   %-5s %d da propria casa, %d da media geral" % (casa, *por_casa[casa]))

    if not gravar:
        print("\nprimeiras linhas:")
        for r, casa, item, media, de in preencher[:8]:
            print("   L%-4d %-5s %-32s R$ %9.4f   <- %s" % (r, casa, item, media, de))
        print("\n(nada foi gravado — rode com --gravar)")
        return 0

    G.BACKUPS.mkdir(exist_ok=True)
    carimbo = datetime.now().strftime("%Y-%m-%d_%H%M%S")
    if not G.copiar_backup(destino, G.BACKUPS / ("cadastro_precos_%s.xlsx" % carimbo)):
        print("\n[erro] nao consegui fazer o backup — nao gravei nada.")
        return 1
    for r, _casa, _item, media, _de in preencher:
        cel = ws.cell(r, 6)
        cel.value = media
        cel.number_format = G.FMT_MOEDA
    wb.save(destino)
    print("\n%d valor(es) preenchido(s). backup em"
          "\nbackups/cadastro_precos_%s.xlsx" % (len(preencher), carimbo))
    return 0


if __name__ == "__main__":
    sys.exit(main())
